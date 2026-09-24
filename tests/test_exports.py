from io import BytesIO
from types import SimpleNamespace
from openpyxl import load_workbook
from radar.export.csv_export import export_csv
from radar.export.html_export import export_html
from radar.export.xlsx_export import export_xlsx


def sample():
    service = SimpleNamespace(service_name="Почта", host="mail.example", port=443, owner="owner")
    result = SimpleNamespace(status="OK", days_left=90, risk_score=10)
    return [(service, result)]


def test_all_exports():
    assert export_csv(sample()).startswith("\ufeff")
    assert "<table>" in export_html(sample())
    sheet = load_workbook(BytesIO(export_xlsx(sample()))).active
    assert sheet.max_row == 2
    assert sheet.auto_filter.ref == "A1:Y2"
