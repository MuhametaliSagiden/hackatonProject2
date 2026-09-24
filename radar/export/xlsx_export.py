from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .csv_export import COLUMNS, extract_row_values

STATUS_COLORS = {
    "OK": ("D4EDDA", "155724"),
    "Information": ("D1ECF1", "0C5460"),
    "Warning": ("FFF3CD", "856404"),
    "Critical": ("FFE5D0", "A84200"),
    "Expired": ("F8D7DA", "721C24"),
    "Unreachable": ("E2E3E5", "383D41"),
}


def export_xlsx(rows: list[tuple[Any, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Сертификаты"

    # Header row
    ws.append(COLUMNS)
    bold_font = Font(bold=True)
    for col_num in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = bold_font

    status_col_index = COLUMNS.index("Статус") + 1

    # Data rows
    for r_idx, (service, result) in enumerate(rows, start=2):
        row_values = extract_row_values(service, result)
        ws.append(row_values)

        status = getattr(result, "status", "")
        if status in STATUS_COLORS:
            bg, fg = STATUS_COLORS[status]
            status_cell = ws.cell(row=r_idx, column=status_col_index)
            status_cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
            status_cell.font = Font(color=fg, bold=True)

    max_row = max(ws.max_row, 1)
    end_col_letter = ws.cell(row=1, column=len(COLUMNS)).column_letter
    ws.auto_filter.ref = f"A1:{end_col_letter}{max_row}"
    ws.freeze_panes = "A2"

    for column in ws.columns:
        col_letter = column[0].column_letter
        max_len = max(len(str(cell.value or "")) for cell in column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    data = BytesIO()
    wb.save(data)
    return data.getvalue()
